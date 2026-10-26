"""Excel版(マクロなし)の生成。

「入力」シートに物件情報を入れ、工事タイプ・階数・構造を選ぶだけで、
「見積書」シートの内訳が数式で自動的に組み上がる。単価は業者が「見積書」の黄色欄に入力する。

実行: python excel/build_excel.py  ->  excel/見積依頼書ジェネレーター.xlsx
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
catalog = json.loads((HERE.parent / "backend" / "catalog.json").read_text(encoding="utf-8"))
cats = catalog["categories"]
presets = catalog["presets"]
NC = len(cats)

UI = "Yu Gothic UI"
MIN = "ＭＳ 明朝"
THIN, HAIR = Side(style="thin"), Side(style="hair")
GREEN = "17844A"
F_HEAD = PatternFill("solid", fgColor="123524")
F_LABEL = PatternFill("solid", fgColor="EEF7F1")
F_INPUT = PatternFill("solid", fgColor="FFF2CC")
F_GREY = PatternFill("solid", fgColor="D9D9D9")


def font(name=UI, size=11, bold=False, color="17261E"):
    return Font(name=name, size=size, bold=bold, color=color)


wb = Workbook()
ws_help = wb.active
ws_help.title = "使い方"
ws_in = wb.create_sheet("入力")
ws_out = wb.create_sheet("見積書")
ws_m = wb.create_sheet("マスタ")
ws_s = wb.create_sheet("設定")

# ======================= 設定 =======================
S = ws_s
S["A1"], S["B1"] = "構造", "キー"
for i, (n, k) in enumerate([("木造", "wood"), ("鉄骨造", "steel"), ("RC造", "rc")], 2):
    S.cell(i, 1, n); S.cell(i, 2, k)
S["D1"], S["E1"] = "階数", "数値"
for i, (n, k) in enumerate([("平屋", 1), ("2階建て", 2), ("3階建て", 3)], 2):
    S.cell(i, 4, n); S.cell(i, 5, k)
S["G1"] = "工事タイプ"
for i, p in enumerate(presets, 2):
    S.cell(i, 7, p["name"])

S["A8"], S["B8"] = "基準キー", "値（入力シートから自動）"
# 入力シートのセル位置（レイアウトを変えるときはここだけ直す）
R = dict(cust="B5", work="B6", addr="B7", company="B8", vendor="B9", period="B10", due="B11", remarks="B12",
         struct="B16", floors="B17", type="B18",
         f1="B21", f2="B22", f3="B23", roof="B24", wall="B25", scaf="B26", toilet="B27", cloth="B28", area_show="B29")
CAT0 = 33  # 入力シートの工事一覧の先頭行
q = lambda k: f"'入力'!{R[k]}"
FLOORS, SKEY, TOTAL, TSUBO, SCAF, PIDX = ("'設定'!$K$%d" % n for n in range(2, 8))

S["J1"], S["K1"] = "自動計算", "値"
S["J2"], S["K2"] = "階数(数値)", f"=MATCH({q('floors')},$D$2:$D$4,0)"
S["J3"], S["K3"] = "構造キー", f"=INDEX($B$2:$B$4,MATCH({q('struct')},$A$2:$A$4,0))"
AREA_RNG = f"'入力'!{R['f1']}:{R['f3']}"
S["J4"], S["K4"] = "延床面積(㎡)", (
    f'=IF(AND({q("f1")}="",IF(K2>=2,{q("f2")},"")="",IF(K2>=3,{q("f3")},"")=""),"",'
    f'N({q("f1")})+IF(K2>=2,N({q("f2")}),0)+IF(K2>=3,N({q("f3")}),0))')
S["J5"], S["K5"] = "延床面積(坪)", '=IF(K4="","",ROUND(K4/3.30579,2))'
S["J6"], S["K6"] = "足場面積", f'=IF({q("scaf")}<>"",{q("scaf")},IF({q("wall")}<>"",{q("wall")},""))'
S["J7"], S["K7"] = "工事タイプ番号", f"=MATCH({q('type')},$G$2:$G${1 + len(presets)},0)"
S.column_dimensions["J"].width = 16; S.column_dimensions["K"].width = 14
basis = [
    ("floor1_area", f'=IF({q("f1")}="","",{q("f1")})'),
    ("floor2_area", f'=IF(AND(K2>=2,{q("f2")}<>""),{q("f2")},"")'),
    ("floor3_area", f'=IF(AND(K2>=3,{q("f3")}<>""),{q("f3")},"")'),
    ("total_area", "=K4"),
    ("tsubo", "=K5"),
    ("roof_area", f'=IF({q("roof")}="","",{q("roof")})'),
    ("wall_area", f'=IF({q("wall")}="","",{q("wall")})'),
    ("scaffold_area", "=K6"),
    ("toilet_count", f'=IF({q("toilet")}="","",{q("toilet")})'),
    ("cloth_len", f'=IF({q("cloth")}="","",{q("cloth")})'),
]
for i, (k, f) in enumerate(basis, 9):
    S.cell(i, 1, k); S.cell(i, 2, f)
B_LAST = 8 + len(basis)

C0 = 23  # 工事テーブルの先頭行
S.cell(C0 - 1, 1, "工事ID"); S.cell(C0 - 1, 2, "工事名")
for j, p in enumerate(presets):
    S.cell(C0 - 1, 3 + j, p["name"])
PC1, PC2 = "C", chr(ord("C") + len(presets) - 1)
S.cell(C0 - 1, 7, "採用"); S.cell(C0 - 1, 8, "採用連番")
for i, c in enumerate(cats):
    r = C0 + i
    S.cell(r, 1, c["id"]); S.cell(r, 2, c["name"])
    for j, p in enumerate(presets):
        S.cell(r, 3 + j, 1 if c["id"] in p["ids"] else 0)
    ir = CAT0 + i  # 入力シート側の行
    S.cell(r, 7, f'=IF(\'入力\'!C{ir}="○",1,IF(\'入力\'!C{ir}="×",0,IF(INDEX({PC1}{r}:{PC2}{r},$K$7)=1,1,0)))')
    S.cell(r, 8, f'=IF(G{r}=1,COUNTIF(G${C0}:G{r},1),"")')
C1 = C0 + NC - 1
for c in range(1, 9):
    S.cell(1, c).font = font(bold=True); S.cell(8, c).font = font(bold=True); S.cell(C0 - 1, c).font = font(bold=True)
S.column_dimensions["A"].width = 14; S.column_dimensions["B"].width = 26
for col in "CDEFGH":
    S.column_dimensions[col].width = 16
S.sheet_properties.tabColor = "999999"

# ======================= マスタ =======================
M = ws_m
heads = ["行種別", "工事ID", "名称", "仕様", "既定数量", "単位", "基準キー", "係数", "最低階数", "構造限定", "採用工事", "表示", "数量", "連番"]
for c, h in enumerate(heads, 1):
    cell = M.cell(1, c, h); cell.font = font(bold=True, color="FFFFFF"); cell.fill = F_HEAD
rows = []
for cat in cats:
    rows.append(("T", cat["id"], cat["name"], "", None, "", "", None, None, ""))
    for it in cat["items"]:
        if it.get("group"):
            rows.append(("G", cat["id"], it["name"], "", None, "", "", None, None, ""))
        else:
            rows.append(("I", cat["id"], it["name"], it.get("spec", ""), it.get("qty", 1), it.get("unit", ""),
                         it.get("basis", ""), it.get("factor") if it.get("basis") else None, it.get("minFloors"),
                         ",".join(it.get("structures", []))))
    rows.append(("S", cat["id"], "小　計", "", None, "", "", None, None, ""))
    rows.append(("B", cat["id"], "", "", None, "", "", None, None, ""))
for r, row in enumerate(rows, 2):
    for c, v in enumerate(row, 1):
        if v not in (None, ""):
            M.cell(r, c, v)
    M.cell(r, 11, f"=INDEX('設定'!$G${C0}:$G${C1},B{r})")
    M.cell(r, 12, f'=IF(K{r}<>1,0,IF(OR(A{r}="T",A{r}="S",A{r}="B"),1,IF(AND(OR(I{r}="",I{r}<={FLOORS}),OR(J{r}="",ISNUMBER(SEARCH({SKEY},J{r})))),1,0)))')
    M.cell(r, 13, f'=IF(A{r}<>"I","",IF(G{r}="",E{r},IFERROR(IF(VLOOKUP(G{r},\'設定\'!$A$9:$B${B_LAST},2,FALSE)="",E{r},ROUND(VLOOKUP(G{r},\'設定\'!$A$9:$B${B_LAST},2,FALSE)*H{r},2)),E{r})))')
    M.cell(r, 14, f'=IF(L{r}=1,COUNTIF(L$2:L{r},1),"")')
MLAST = len(rows) + 1
for col, w in zip("ABCDEFGHIJKLMN", [8, 8, 40, 30, 10, 8, 14, 8, 10, 12, 10, 8, 10, 8]):
    M.column_dimensions[col].width = w
M.freeze_panes = "A2"
M.sheet_properties.tabColor = "999999"

# ======================= 入力 =======================
I = ws_in
I.sheet_view.showGridLines = False
for col, w in zip("ABCD", [30, 30, 22, 62]):
    I.column_dimensions[col].width = w
F_BAND = PatternFill("solid", fgColor=GREEN)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
I["A1"] = "見積依頼書ジェネレーター"; I["A1"].font = font(size=20, bold=True, color="123524")
I["A2"] = "STEP 1 → 5 の順に、黄色のセルへ入力・選択してください。（入力済みの内容は例です。上書きしてください）"
I["A2"].font = font(size=11, color="66786D")


def band(r, text):
    for c in range(1, 5):
        I.cell(r, c).fill = F_BAND
    I.cell(r, 1, text).font = font(size=13, bold=True, color="FFFFFF")
    I.row_dimensions[r].height = 24


def lab(r, text, hint="", unit=""):
    c = I.cell(r, 1, text); c.font = font(bold=True); c.fill = F_LABEL; c.border = BOX
    if unit:
        I.cell(r, 3, unit).font = font(size=10, color="66786D")
    if hint:
        h = I.cell(r, 4, hint); h.font = font(size=10, color="66786D"); h.alignment = Alignment(vertical="center", wrap_text=True)


def box(addr, value=None, fmt=None, align="left"):
    c = I[addr]; c.value = value; c.fill = F_INPUT; c.font = font(); c.border = BOX
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if fmt:
        c.number_format = fmt


band(4, "STEP 1　物件の基本情報")
for key, t, hint, v in [
    ("cust", "お客様名", "例: 山田太郎　（「様邸」は自動で付きます）", None),
    ("work", "工事名", "例: リノベーション工事", "リノベーション工事"),
    ("addr", "物件所在地", "", None),
    ("company", "依頼元（自社名）", "見積書の右上に表示されます", None),
    ("vendor", "提出先（業者名）", "「御中」は自動で付きます。業者ごとに変えて見積書を作ります", None),
    ("period", "工事希望時期", "例: 2026年12月〜", None),
    ("due", "見積提出期限", "日付を入力（例: 2026/11/30）", None),
    ("remarks", "備考・特記事項", "搬入経路、工事時間の制限など、業者に伝えたいこと", None),
]:
    addr = R[key]; lab(int(addr[1:]), t, hint)
    box(addr, v, "yyyy/m/d" if key == "due" else None)
I.row_dimensions[12].height = 46

band(14, "STEP 2　建物の条件を選ぶ")
I["A15"] = "▼ が出るセルは、一覧から選びます。"; I["A15"].font = font(size=10, color="66786D")
lab(16, "構造", "木造以外では、木工事の木材項目が外れ、下地の項目に切り替わります"); box("B16", "木造")
lab(17, "階数", "選んだ階数に合わせて、2階・3階の項目だけが出ます"); box("B17", "2階建て")
lab(18, "工事タイプ", "選ぶと、含める工事がSTEP 4の表で自動で決まります"); box("B18", presets[0]["name"])
for r in (16, 17, 18):
    I.row_dimensions[r].height = 24

band(20, "STEP 3　面積などの数値を入力")
for key, t, unit, hint, v in [
    ("f1", "1階床面積", "㎡", "解体・木工事・防蟻・養生清掃などの数量に使います", 63.6),
    ("f2", "2階床面積", "㎡", "平屋のときは使いません", 33),
    ("f3", "3階床面積", "㎡", "3階建てのときだけ使います", None),
    ("roof", "屋根面積", "㎡", "屋根板金工事の屋根葺き・防水シートに使います", 83),
    ("wall", "外壁面積", "㎡", "外壁の角波工事に使います", 123.3),
    ("scaf", "足場面積", "㎡", "仮設足場・メッシュシートに使います（空欄なら外壁面積）", 360),
    ("toilet", "トイレの数", "ヶ所", "トイレ解体の数量に使います", 2),
    ("cloth", "クロス貼りの延長", "m", "クロス貼りの数量に使います", 670),
]:
    addr = R[key]; lab(int(addr[1:]), t, hint, unit); box(addr, v, "General", "right")
lab(29, "→ 延床面積（自動計算）")
I["B29"] = f'=IF({TOTAL}="","（未入力）",{TOTAL}&" ㎡（約 "&{TSUBO}&" 坪）")'
I["B29"].font = font(bold=True, color=GREEN); I["B29"].border = BOX
I["D29"] = "選んだ階数の床面積の合計です"; I["D29"].font = font(size=10, color="66786D")

band(31, "STEP 4　含める工事を確認する")
for c, h in zip("ABCD", ["工事", "タイプでの初期値", "手動で変更（○=含める / ×=外す）", "結果（項目数）"]):
    I[f"{c}32"] = h; I[f"{c}32"].font = font(bold=True, color="FFFFFF"); I[f"{c}32"].fill = F_HEAD
    I[f"{c}32"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
I.row_dimensions[32].height = 34
for i, cat in enumerate(cats):
    r = CAT0 + i; sr = C0 + i
    I.cell(r, 1, cat["name"]).font = font(bold=True)
    I.cell(r, 2, f"=IF(INDEX('設定'!${PC1}${sr}:${PC2}${sr},{PIDX})=1,\"含める\",\"—\")").font = font()
    I.cell(r, 3).fill = F_INPUT
    I.cell(r, 4, f"=IF('設定'!G{sr}=1,\"採用（\"&COUNTIFS('マスタ'!$B$2:$B${MLAST},{cat['id']},'マスタ'!$A$2:$A${MLAST},\"I\",'マスタ'!$L$2:$L${MLAST},1)&\"項目）\",\"除外\")").font = font(bold=True)
    for c in range(1, 5):
        I.cell(r, c).border = BOX
    for c in (2, 3):
        I.cell(r, c).alignment = Alignment(horizontal="center")
    I.row_dimensions[r].height = 21
last = CAT0 + NC - 1
I.cell(last + 1, 1, "現在の内訳").font = font(bold=True, color=GREEN)
I.cell(last + 1, 2, f"=COUNTIF('設定'!$G${C0}:$G${C1},1)&\"工事 / \"&COUNTIFS('マスタ'!$A$2:$A${MLAST},\"I\",'マスタ'!$L$2:$L${MLAST},1)&\"項目\"").font = font(bold=True, color=GREEN)
I.conditional_formatting.add(f"D{CAT0}:D{last}", FormulaRule(formula=[f'LEFT($D{CAT0},2)="採用"'], font=Font(color=GREEN, bold=True), fill=PatternFill("solid", bgColor="EEF7F1", fgColor="EEF7F1")))
I.conditional_formatting.add(f"D{CAT0}:D{last}", FormulaRule(formula=[f'$D{CAT0}="除外"'], font=Font(color="999999")))

s5 = last + 3
band(s5, "STEP 5　見積書を確認して、業者へ送る")
I.cell(s5 + 1, 1, '=HYPERLINK("#\'見積書\'!A1","▶ 見積書シートを開く")').font = Font(name=UI, size=12, bold=True, color="0563C1", underline="single")
for k, t in enumerate([
    "見積書シートに、内訳が自動で並んでいます。内容を確認してください。",
    "業者へ送るときは、このファイルを別名で保存して送ります。業者は「見積書」の黄色い欄（単価・摘要）だけ入力します。",
    "送ったあとは、このファイルの入力内容を変えないでください（行の位置がずれます）。",
], start=2):
    I.cell(s5 + k, 1, "・" + t).font = font(size=11)


def dvp(f, title, msg):
    return DataValidation(type="list", formula1=f, allow_blank=False, showInputMessage=True, promptTitle=title, prompt=msg)


dvs = [
    (dvp("'設定'!$A$2:$A$4", "構造", "一覧から選びます"), R["struct"]),
    (dvp("'設定'!$D$2:$D$4", "階数", "一覧から選びます"), R["floors"]),
    (dvp("'設定'!$G$2:$G$%d" % (1 + len(presets)), "工事タイプ", "一覧から選びます"), R["type"]),
    (DataValidation(type="list", formula1='"○,×"', allow_blank=True, showInputMessage=True, promptTitle="手動で変更",
                    prompt="含めるなら ○、外すなら × を選びます。空欄ならタイプの初期値のままです"), f"C{CAT0}:C{last}"),
]
for d, rng in dvs:
    I.add_data_validation(d); d.add(rng)
I.sheet_properties.tabColor = GREEN
I.freeze_panes = "A4"


# ======================= 見積書 =======================
O = ws_out
for col, w in zip("ABCDEFGHIJK", [6, 38, 34, 9, 2, 7, 14, 2, 16, 2, 28]):
    O.column_dimensions[col].width = w
D0 = 24
DN = len(rows) + 6
DL = D0 + DN - 1
SUMR0 = 5

O["B1"] = f'=IF({q("cust")}="","お名前",{q("cust")})&" 様邸"&{q("work")}&" 見積書"'
O["B1"].font = font(MIN, 14, True)
O["B2"] = f'=IF({q("vendor")}="","",{q("vendor")}&" 御中")'; O["B2"].font = font(MIN)
O.merge_cells("G2:K2")
O["G2"] = f'={q("company")}&""'; O["G2"].font = font(MIN); O["G2"].alignment = Alignment(horizontal="right")
for c in range(1, 12):
    cell = O.cell(3, c); cell.fill = F_GREY
    cell.border = Border(top=THIN, bottom=THIN, left=THIN if c in (1, 2, 3, 4, 6, 7, 9, 11) else None, right=THIN if c == 11 else None)
for c, t in {1: "No.", 2: "名　　　称", 3: "仕　　　様", 4: "数量", 6: "単位", 7: "単　価", 9: "金　額", 11: "摘　要"}.items():
    O.cell(3, c, t).font = font(MIN, bold=True); O.cell(3, c).alignment = Alignment(horizontal="center")
for a, b in ((4, 5), (7, 8), (9, 10)):
    O.merge_cells(start_row=3, start_column=a, end_row=3, end_column=b)
O["A4"] = "（内　訳　明　細）"; O["A4"].font = font(MIN)

SUMN = NC
for j in range(1, SUMN + 1):
    r = SUMR0 + j - 1
    O.cell(r, 14, f"=IFERROR(MATCH({j},'設定'!$H${C0}:$H${C1},0),\"\")")
    O.cell(r, 1, f'=IF(N{r}="","",{j})')
    O.cell(r, 2, f"=IF(N{r}=\"\",\"\",INDEX('設定'!$B${C0}:$B${C1},N{r}))")
    O.cell(r, 4, f'=IF(N{r}="","",1)')
    O.cell(r, 6, f'=IF(N{r}="","","式")')
    O.cell(r, 9, f'=IF(N{r}="","",SUMIFS(I${D0}:I${DL},N${D0}:N${DL},N{r},M${D0}:M${DL},"S"))')
    O.cell(r, 9).number_format = "#,##0"
    for c in (1, 2, 4, 6, 9):
        O.cell(r, c).font = font(MIN)
    for c in range(1, 12):
        O.cell(r, c).border = Border(left=THIN if c in (1, 2, 3, 4, 6, 7, 9, 11) else None, right=THIN if c in (1, 2, 3, 5, 6, 8, 10, 11) else None, top=HAIR, bottom=HAIR)
ADJ = SUMR0 + SUMN + 1
TOT = ADJ + 2
O.cell(ADJ, 2, "調整費").font = font(MIN)
O.cell(ADJ, 9).fill = F_INPUT; O.cell(ADJ, 9).number_format = "#,##0;[Red]-#,##0"
O.cell(TOT, 2, "合　計").font = font(MIN, bold=True)
O.cell(TOT, 9, f"=SUM(I{SUMR0}:I{SUMR0 + SUMN - 1})+N(I{ADJ})").font = font(MIN, bold=True)
O.cell(TOT, 9).number_format = "#,##0"
for c in range(1, 12):
    O.cell(TOT, c).border = Border(top=THIN, bottom=THIN, left=THIN if c in (1, 2, 3, 4, 6, 7, 9, 11) else None, right=THIN if c == 11 else None)

O["B21"] = (f'="構造: "&{q("struct")}&" / "&{q("floors")}&IF({TOTAL}="","",", 延床 "&{TOTAL}&"㎡（"&{TSUBO}&"坪）")'
            f'&IF({q("due")}="","", " / 提出期限 "&TEXT({q("due")},"yyyy/m/d"))')
O["B22"] = f'=IF({q("remarks")}="","","備考: "&{q("remarks")})'
O["B21"].font = font(MIN, 10); O["B22"].font = font(MIN, 10)
O.merge_cells("G1:K1"); O["G1"] = "※ 黄色の欄（単価・摘要）をご記入ください"
O["G1"].font = font(MIN, 10, True, "C2410C"); O["G1"].alignment = Alignment(horizontal="right")
O["A23"] = "（内　訳　明　細）"; O["A23"].font = font(MIN)

MR = lambda col: f"'マスタ'!${col}$2:${col}${MLAST}"
for r in range(D0, DL + 1):
    k = r - D0 + 1
    O.cell(r, 12, f'=IFERROR(MATCH({k},{MR("N")},0),"")')
    O.cell(r, 13, f'=IF(L{r}="","",INDEX({MR("A")},L{r}))')
    O.cell(r, 14, f'=IF(L{r}="","",INDEX({MR("B")},L{r}))')
    O.cell(r, 1, f"=IF(M{r}=\"T\",INDEX('設定'!$H${C0}:$H${C1},N{r}),\"\")")
    O.cell(r, 2, f'=IF(OR(L{r}="",M{r}="B"),"",INDEX({MR("C")},L{r}))')
    O.cell(r, 3, f'=IF(M{r}="I",INDEX({MR("D")},L{r})&"","")')
    O.cell(r, 4, f'=IF(M{r}="I",INDEX({MR("M")},L{r}),"")')
    O.cell(r, 6, f'=IF(M{r}="I",INDEX({MR("F")},L{r})&"","")')
    if r == D0:
        O.cell(r, 9, f'=IF(AND(M{r}="I",G{r}<>"",D{r}<>""),ROUND(D{r}*G{r},0),"")')
    else:
        O.cell(r, 9, f'=IF(M{r}="I",IF(OR(G{r}="",D{r}=""),"",ROUND(D{r}*G{r},0)),IF(M{r}="S",SUMIFS(I${D0}:I{r - 1},N${D0}:N{r - 1},N{r},M${D0}:M{r - 1},"I"),""))')
    O.cell(r, 9).number_format = "#,##0"
    O.cell(r, 7).number_format = "#,##0"
    O.cell(r, 4).alignment = Alignment(horizontal="right")
    O.cell(r, 6).alignment = Alignment(horizontal="center")
    for c in (1, 2, 3, 4, 6, 7, 9, 11):
        O.cell(r, c).font = font(MIN)

rng_all = f"A{D0}:K{DL}"
def cf(sqref, formula, **kw):
    O.conditional_formatting.add(sqref, FormulaRule(formula=[formula], **kw))
left = f"A{D0}:D{DL} F{D0}:G{DL} I{D0}:I{DL} K{D0}:K{DL}"
right = f"E{D0}:E{DL} H{D0}:H{DL} J{D0}:K{DL}"
cf(left, f'$L{D0}<>""', border=Border(left=THIN, top=HAIR, bottom=HAIR))
cf(f"B{D0}:C{DL}", f'$L{D0}<>""', border=Border(top=HAIR, bottom=HAIR))
cf(right, f'$L{D0}<>""', border=Border(right=THIN, top=HAIR, bottom=HAIR))
cf(f"A{D0}:K{DL}", f'$M{D0}="T"', font=Font(bold=True))
cf(f"A{D0}:K{DL}", f'$M{D0}="G"', font=Font(bold=True))
cf(f"A{D0}:K{DL}", f'$M{D0}="S"', font=Font(bold=True), border=Border(top=THIN, bottom=THIN, left=THIN, right=THIN))
cf(f"G{D0}:G{DL} K{D0}:K{DL}", f'$M{D0}="I"', fill=PatternFill("solid", bgColor="FFF2CC", fgColor="FFF2CC"))
O.column_dimensions.group("L", "N", hidden=True)
O.freeze_panes = "A4"
O.page_setup.orientation = "portrait"; O.page_setup.paperSize = O.PAPERSIZE_A4
O.page_setup.fitToWidth = 1; O.page_setup.fitToHeight = 0
O.sheet_properties.pageSetUpPr.fitToPage = True
O.print_title_rows = "3:3"
O.sheet_properties.tabColor = "C2410C"

# ======================= 使い方 =======================
H = ws_help
H.sheet_view.showGridLines = False
H.column_dimensions["A"].width = 6; H.column_dimensions["B"].width = 34; H.column_dimensions["C"].width = 80
H["A1"] = "見積依頼書ジェネレーター（Excel版）"; H["A1"].font = font(size=20, bold=True, color="123524")
H["A2"] = "物件の情報を入れると、業者に渡す見積依頼書（単価欄は空欄）が自動でできあがります。マクロは使っていません。"
H["A2"].font = font(size=11, color="66786D")
for c, t in zip("ABC", ["", "やること", "内容"]):
    H[f"{c}4"] = t; H[f"{c}4"].font = font(bold=True, color="FFFFFF"); H[f"{c}4"].fill = F_HEAD
steps = [
    ("1", "「入力」シートを開く", "STEP 1〜4 の黄色いセルに、お客様名・構造・階数・面積などを入れ、工事タイプを選びます。"),
    ("2", "含める工事を確認する", "STEP 4 の表で、採用される工事が分かります。変えたいときだけ「手動で変更」に ○ か × を入れます。"),
    ("3", "「見積書」シートで内訳を見る", "選んだ内容に合わせて、工事ごとの内訳と数量が自動で並びます。"),
    ("4", "業者に送る", "ファイルを別名で保存して送ります。業者は「見積書」の黄色い欄（単価・摘要）に入力します。"),
    ("5", "返ってきた見積書を見る", "業者が単価を入れると、金額・小計・合計が自動で計算されます。"),
]
for r, (n, a, b) in enumerate(steps, 5):
    H.cell(r, 1, n).font = font(size=16, bold=True, color=GREEN); H.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
    H.cell(r, 2, a).font = font(bold=True); H.cell(r, 3, b).font = font()
    for c in (2, 3):
        H.cell(r, c).alignment = Alignment(wrap_text=True, vertical="center")
    for c in range(1, 4):
        H.cell(r, c).border = Border(bottom=HAIR)
    H.row_dimensions[r].height = 34

H["A11"] = "シートの役割"; H["A11"].font = font(size=13, bold=True, color=GREEN)
for r, (a, b) in enumerate([
    ("入力", "自分で入力・選択するシートです。ふだん触るのはここだけです。"),
    ("見積書", "業者に渡すシートです。内訳は自動で並びます。単価と摘要だけ業者が入力します。"),
    ("マスタ", "全項目の一覧です。項目の名前や単位を直したいときに使います。灰色のタブは、ふだん触らなくて大丈夫です。"),
], 12):
    H.cell(r, 2, a).font = font(bold=True); H.cell(r, 3, b).font = font()
    H.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="center"); H.row_dimensions[r].height = 30

H["A16"] = "気をつけること"; H["A16"].font = font(size=13, bold=True, color=GREEN)
for r, t in enumerate([
    "業者に送ったあとは、自社のファイルの「入力」シートを変えないでください。内訳の行の位置がずれて、業者の単価とずれます。",
    "数式を含まない見積書にしたいときは、「見積書」シートを全選択してコピーし、新しいブックに「値と元の書式設定」で貼り付けます。",
    "項目を足したいときは、マスタの小計行（行種別 S）の上に行を挿入し、K〜N列の数式を上の行からコピーしてください。",
    "「設定」シートは非表示です。工事タイプの中身を変えるときは、シートのタブを右クリック →「再表示」で開けます。",
], 17):
    H.cell(r, 2, "・" + t).font = font(); H.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    H.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top"); H.row_dimensions[r].height = 34
H["A22"] = "元のフォーマット（リフォーム・リノベーション工事見積書）の体裁に合わせています。"
H["A22"].font = font(size=10, color="66786D")
H.sheet_properties.tabColor = "123524"

ws_s.sheet_state = "hidden"


wb.active = 0
out = HERE / "見積依頼書ジェネレーター.xlsx"
wb.save(out)
print("saved", out, "master rows", len(rows), "detail rows", DN)
