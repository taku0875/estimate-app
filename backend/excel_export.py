"""見積依頼書(Excel)の生成。元フォーマット「リフォーム・リノベーション工事見積書」の体裁を踏襲する。

単価は業者が入力する前提のため空欄にし、金額・小計・合計は数式にしてある。
ブラウザ版(web/export_xlsx.js)も同じレイアウトで作っているので、変えるときは両方そろえる。

「見積書」シートの並び:
    ヘッダ → 大項目の一覧(金額は各小計を参照) → 調整費 → 合計 → 大項目ごとの内訳と小計
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "ＭＳ 明朝"
THIN = Side(style="thin")
HAIR = Side(style="hair")
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # 業者の入力欄(黄色)
MONEY = "#,##0"

# ---- 列 ----
# E・H・J 列は元フォーマットにある細い余白列。
COL_NO, COL_NAME, COL_SPEC, COL_QTY, COL_UNIT, COL_PRICE, COL_AMOUNT, COL_NOTE = 1, 2, 3, 4, 6, 7, 9, 11
LAST_COL = COL_NOTE
AMOUNT = "I"  # 金額列(数式で参照する)
COL_WIDTHS = {"A": 6, "B": 38, "C": 34, "D": 9, "E": 2, "F": 7, "G": 14, "H": 2, "I": 16, "J": 2, "K": 28}
# 縦罫線を引く列(列の左側に線を引く)。余白列は直前の列とひと続きに見せる。
RULED_COLS = (COL_NO, COL_NAME, COL_SPEC, COL_QTY, COL_UNIT, COL_PRICE, COL_AMOUNT, COL_NOTE)

FLOOR_LABELS = {1: "平屋", 2: "2階建て", 3: "3階建て"}


def _font(size=11, bold=False):
    return Font(name=FONT, size=size, bold=bold)


def _qty(v):
    """数量の表示用。整数ならintに、それ以外は小数2桁に丸める。"""
    if v is None or v == "":
        return None
    return int(v) if float(v) == int(v) else round(float(v), 2)


def _with_unit(v, unit):
    if v in (None, ""):
        return None
    return f"{_qty(v)} {unit}"


def build_workbook(data: dict) -> BytesIO:
    wb = Workbook()
    _write_estimate_sheet(wb.active, data["property"], data["categories"])
    _write_property_sheet(wb.create_sheet("物件情報"), data["property"])
    _neutralize_formulas(wb)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- 見積書シート

def _write_estimate_sheet(ws, prop, cats):
    ws.title = "見積書"
    for col, w in COL_WIDTHS.items():
        ws.column_dimensions[col].width = w

    def put(row, col, value=None, *, bold=False, fmt=None, fill=None, align=None):
        cell = ws.cell(row, col, value)
        cell.font = _font(bold=bold)
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = fill
        if align:
            cell.alignment = Alignment(horizontal=align)
        return cell

    def rule(row, line=HAIR):
        """1行分の罫線。line=THIN で上下を実線にする(ヘッダ・合計・小計)。"""
        for c in range(1, LAST_COL + 1):
            ws.cell(row, c).border = Border(
                left=THIN if c in RULED_COLS else None,
                right=THIN if c == LAST_COL or c + 1 in RULED_COLS else None,
                top=line, bottom=line,
            )

    # ---- タイトル・宛先 ----
    title = f"{prop.get('customer_name') or 'お名前'} 様邸{prop.get('work_title') or 'リノベーション工事'} 見積書"
    ws.cell(1, COL_NAME, title).font = _font(14, True)
    ws.row_dimensions[1].height = 24
    put(2, COL_NAME, f"{prop['vendor_name']} 御中" if prop.get("vendor_name") else "")
    ws.merge_cells(start_row=2, start_column=COL_PRICE, end_row=2, end_column=LAST_COL)
    put(2, COL_PRICE, prop.get("company_name") or "", align="right")

    # ---- 見出し行 ----
    head_row = 3
    rule(head_row, THIN)
    headings = {COL_NO: "No.", COL_NAME: "名　　　称", COL_SPEC: "仕　　　様", COL_QTY: "数量",
                COL_UNIT: "単位", COL_PRICE: "単　価", COL_AMOUNT: "金　額", COL_NOTE: "摘　要"}
    for c in range(1, LAST_COL + 1):
        ws.cell(head_row, c).fill = HEADER_FILL
    for c, text in headings.items():
        put(head_row, c, text, bold=True).alignment = Alignment(horizontal="center", vertical="center")
    for c in (COL_QTY, COL_PRICE, COL_AMOUNT):  # 右隣の余白列と結合
        ws.merge_cells(start_row=head_row, start_column=c, end_row=head_row, end_column=c + 1)

    # ---- 大項目の一覧(金額はあとで各小計への参照を入れる) ----
    row = head_row + 1
    put(row, COL_NO, "（内　訳　明　細）")
    row += 1
    summary_rows = []
    for i, cat in enumerate(cats, 1):
        put(row, COL_NO, i)
        put(row, COL_NAME, cat["name"])
        put(row, COL_QTY, 1)
        put(row, COL_UNIT, "式")
        rule(row)
        summary_rows.append(row)
        row += 1

    adj_row = row + 1
    put(adj_row, COL_NAME, "調整費")
    put(adj_row, COL_AMOUNT, fmt="#,##0;[Red]-#,##0", fill=INPUT_FILL)
    rule(adj_row)

    total_row = adj_row + 2
    put(total_row, COL_NAME, "合　計", bold=True)
    rule(total_row, THIN)

    # ---- 大項目ごとの内訳 ----
    row = total_row + 2
    subtotal_rows = []
    for i, cat in enumerate(cats, 1):
        put(row, COL_NO, "（内　訳　明　細）")
        row += 1
        put(row, COL_NO, i, bold=True)
        put(row, COL_NAME, cat["name"], bold=True)
        rule(row)
        row += 1

        first = row
        for it in cat["items"]:
            _write_item_row(put, row, it)
            rule(row)
            row += 1

        put(row, COL_NAME, "小　計", bold=True)
        # 項目が0件のとき SUM(I{first}:I{first-1}) は小計セル自身を含む循環参照になるので 0 にする
        subtotal = f"=SUM({AMOUNT}{first}:{AMOUNT}{row - 1})" if row > first else 0
        put(row, COL_AMOUNT, subtotal, bold=True, fmt=MONEY)
        rule(row, THIN)
        subtotal_rows.append(row)
        row += 2

    # 一覧の金額 = 各小計、合計 = 一覧の合計 + 調整費
    for summary_row, subtotal_row in zip(summary_rows, subtotal_rows):
        put(summary_row, COL_AMOUNT, f"={AMOUNT}{subtotal_row}", fmt=MONEY)
    total = (f"=SUM({AMOUNT}{summary_rows[0]}:{AMOUNT}{summary_rows[-1]})+N({AMOUNT}{adj_row})"
             if summary_rows else 0)
    put(total_row, COL_AMOUNT, total, bold=True, fmt=MONEY)

    # ---- 印刷・表示 ----
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{head_row}:{head_row}"


def _write_item_row(put, row, it):
    """内訳の1行。見出し(【 】)は名称だけを太字で書く。"""
    if it.get("group"):
        put(row, COL_NAME, it["name"], bold=True)
        return
    put(row, COL_NAME, it["name"])
    put(row, COL_SPEC, it.get("spec") or None)
    put(row, COL_QTY, _qty(it.get("qty")), fmt="#,##0.##", align="right")
    put(row, COL_UNIT, it.get("unit") or None, align="center")
    put(row, COL_PRICE, fmt=MONEY, fill=INPUT_FILL)  # 単価: 業者が入力
    put(row, COL_AMOUNT, f'=IF(G{row}="","",ROUND(D{row}*G{row},0))', fmt=MONEY)
    put(row, COL_NOTE, it.get("note") or None, fill=INPUT_FILL)


# ---------------------------------------------------------------- 物件情報シート

def _write_property_sheet(ps, prop):
    ps.column_dimensions["A"].width = 22
    ps.column_dimensions["B"].width = 40
    ps.column_dimensions["C"].width = 10
    ps.cell(1, 1, "物件情報（見積条件）").font = _font(14, True)
    rows = [
        ("お客様名", prop.get("customer_name")),
        ("工事名", prop.get("work_title")),
        ("物件所在地", prop.get("address")),
        ("構造・工法", prop.get("structure")),
        ("依頼元", prop.get("company_name")),
        ("提出先（業者）", prop.get("vendor_name")),
        ("見積提出期限", prop.get("due_date")),
        ("工事希望時期", prop.get("work_period")),
        ("延床面積", _with_unit(prop.get("total_area"), "㎡")),
        ("延床面積（坪）", _with_unit(prop.get("tsubo"), "坪")),
        ("階数", FLOOR_LABELS.get(prop.get("floors"))),
        ("1階床面積", _with_unit(prop.get("floor1_area"), "㎡")),
        ("2階床面積", _with_unit(prop.get("floor2_area"), "㎡")),
        ("3階床面積", _with_unit(prop.get("floor3_area"), "㎡")),
        ("屋根面積", _with_unit(prop.get("roof_area"), "㎡")),
        ("外壁面積", _with_unit(prop.get("wall_area"), "㎡")),
        ("足場面積", _with_unit(prop.get("scaffold_area"), "㎡")),
        ("トイレ数", _with_unit(prop.get("toilet_count"), "ヶ所")),
        ("備考・特記事項", prop.get("remarks")),
    ]
    box = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    for r, (label, value) in enumerate(rows, 3):
        a = ps.cell(r, 1, label)
        a.font = _font(bold=True)
        a.fill = HEADER_FILL
        b = ps.cell(r, 2, value if value not in (None, "") else None)
        b.font = _font()
        b.alignment = Alignment(wrap_text=True, vertical="top")
        a.border = b.border = box
    note_row = len(rows) + 4
    ps.cell(note_row, 1, "※ 黄色のセル（単価・摘要）をご記入ください。金額・小計・合計は自動計算されます。").font = _font()


# ---------------------------------------------------------------- 安全対策

def _neutralize_formulas(wb):
    """利用者の入力(名称・備考など)が「=」で始まると openpyxl は数式として書き込む。
    外部URLを開く数式などを業者側のExcelで実行させないよう、こちらで組んだ金額列以外は文字列に戻す。"""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and not (ws.title == "見積書" and cell.column == COL_AMOUNT):
                    cell.data_type = "s"
